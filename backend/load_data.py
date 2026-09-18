"""Load the AROL dataset (Excel workbook) into a SQLite database.

Run once before starting the backend:  python backend/load_data.py
"""

import hashlib
import sqlite3

import openpyxl

from config import DATASET_XLSX, DATA_DIR, DB_PATH, DEMO_PASSWORD

# Tables the platform needs itself: sessions, chat history and passwords.
PLATFORM_TABLES = """
CREATE TABLE Sessions (
    token TEXT PRIMARY KEY,
    userId TEXT,
    createdAt TEXT
);
CREATE TABLE Messages (
    messageId INTEGER PRIMARY KEY AUTOINCREMENT,
    token TEXT,
    machineId TEXT,
    role TEXT,
    content TEXT,
    agent TEXT,
    sources TEXT,
    createdAt TEXT
);
CREATE TABLE Credentials (
    userId TEXT PRIMARY KEY,
    passwordHash TEXT
);
"""


def create_credentials(conn):
    """Give every account the same demo password, stored as a hash."""
    password_hash = hashlib.sha256(DEMO_PASSWORD.encode()).hexdigest()
    users = conn.execute("SELECT userId FROM Users").fetchall()
    conn.executemany(
        "INSERT INTO Credentials (userId, passwordHash) VALUES (?, ?)",
        [(user[0], password_hash) for user in users],
    )
    return len(users)


def load_sheet(conn, sheet):
    rows = sheet.iter_rows(values_only=True)
    columns = [str(c) for c in next(rows)]
    columns_sql = ", ".join(f'"{c}"' for c in columns)
    placeholders = ", ".join("?" for _ in columns)
    conn.execute(f'CREATE TABLE "{sheet.title}" ({columns_sql})')

    count = 0
    for row in rows:
        if all(v is None for v in row):
            continue
        # Keep empty cells as NULL: some foreign keys are empty in the dataset.
        values = [v if v != "" else None for v in row]
        conn.execute(
            f'INSERT INTO "{sheet.title}" ({columns_sql}) VALUES ({placeholders})',
            values,
        )
        count += 1
    return count


def main():
    DATA_DIR.mkdir(exist_ok=True)
    if DB_PATH.exists():
        DB_PATH.unlink()

    workbook = openpyxl.load_workbook(DATASET_XLSX, read_only=True)
    conn = sqlite3.connect(DB_PATH)
    for sheet in workbook.worksheets:
        count = load_sheet(conn, sheet)
        print(f"{sheet.title}: {count} rows")
    conn.executescript(PLATFORM_TABLES)
    print(f"Credentials: {create_credentials(conn)} accounts (password '{DEMO_PASSWORD}')")
    conn.commit()
    conn.close()
    print(f"Database written to {DB_PATH}")


if __name__ == "__main__":
    main()
